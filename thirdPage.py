import re
import pandas as pd
from docx import Document
from openpyxl import Workbook

# Открываем документ
file_path = 'processed_vpo1.docx'
doc = Document(file_path)

# Шаблоны поиска
patterns = {
    "по графам": re.compile(r"по графам (\d+(?:-\d+|(?:,\s*\d+)*))"),
    "в графах": re.compile(r"в графах (\d+(?:-\d+|(?:,\s*\d+)*))"),
    "по графе": re.compile(r"по графе (\d+)"),
    "в графе": re.compile(r"в графе (\d+)"),
}

# Списки для данных
rows = []
current_section = ""


# Функция для извлечения номера раздела
def extract_section_number(text):
    match = re.match(r"(\S+)(?=\s|$)", text)
    return match.group(1) if match else ""


# Обрабатываем абзацы документа
for para in doc.paragraphs:
    # Определение текущего раздела
    if para.runs[0].bold:
        section_number = extract_section_number(para.text.strip())
        if section_number:
            current_section = section_number
        continue

    # Проверка каждого шаблона
    for key, pattern in patterns.items():
        match = pattern.search(para.text.lower())
        if match:
            numbers = match.group(1)
            # Обработка диапазонов
            if '-' in numbers:
                start, end = map(int, numbers.split('-'))
                numbers_list = list(range(start, end + 1))
            else:
                numbers_list = [num.strip() for num in re.split(r',\s*', numbers)]

            # Сохранение данных
            for number in numbers_list:
                rows.append([current_section, para.text.strip(), number])
            break

# Создание Excel файла
wb = Workbook()
ws = wb.active
ws.title = "Данные"

# Запись заголовков
ws.append(["Раздел", "Текст абзаца", "Число"])

# Запись данных
for row in rows:
    ws.append(row)

# Сохранение файла
output_file = 'output.xlsx'
wb.save(output_file)
print(f"Данные сохранены в файл {output_file}")
